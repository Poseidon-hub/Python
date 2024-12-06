from docx import Document

# Открытие файла
doc = Document("file.docx")

# Чтение таблицы
table = doc.tables[0]
data_dict = {}

# Преобразование данных таблицы в словарь
for i in range(2, 3):
    header = table.cell(0, i).text
    data_dict[header] = {
        "Flash": table.cell(1, i).text,
        "SRAM": table.cell(2, i).text,
        "EEPROM": table.cell(3, i).text
    }

# Вывод в консоль
print(data_dict)
